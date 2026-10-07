"""CVs: uploads converted to Markdown, stored next to the master CV.

All CVs describe the same candidate. The master CV (`config.cv_path`, usually
`cvs/master.md`) is the main one; other `*.md` files in the same folder are reference
CVs (older or role-specific versions). Ranking and drafting read them all. Making a
reference CV the master copies its text into the master file, after backing the old
master up, so the configured path never changes (Docker pins it with JOBSEARCHER_CV).
"""

from __future__ import annotations

import io
import re
import shutil
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from jobsearcher.companies.config import slugify

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
UPLOAD_TYPES = (".pdf", ".docx", ".md", ".markdown", ".txt")


class CvError(ValueError):
    pass


@dataclass
class CvFile:
    name: str  # file stem, also the URL key
    path: Path
    is_master: bool
    size: int
    modified: datetime


def cv_dir(master: Path) -> Path:
    return master.parent


def list_cvs(master: Path) -> list[CvFile]:
    """The master first, then reference CVs by name."""
    folder = cv_dir(master)
    if not folder.is_dir():
        return []
    out = []
    for path in sorted(folder.glob("*.md")):
        stat = path.stat()
        out.append(
            CvFile(
                name=path.stem,
                path=path,
                is_master=path.resolve() == master.resolve(),
                size=stat.st_size,
                modified=datetime.fromtimestamp(stat.st_mtime, UTC),
            )
        )
    out.sort(key=lambda cv: (not cv.is_master, cv.name))
    return out


def ranking_cv(master: Path) -> str | None:
    """Every CV as one text for ranking: the master, then the others as more facts about
    the same candidate (exact copies skipped, e.g. the CV that was made the master).
    None without a master CV."""
    files = list_cvs(master)
    if not files or not files[0].is_master:
        return None
    seen = {files[0].path.read_text().strip()}
    parts = [*seen]
    for cv in files[1:]:
        text = cv.path.read_text().strip()
        if text and text not in seen:
            seen.add(text)
            parts += ["", f"# Other CV: {cv.name} (more facts about the same candidate)", text]
    return "\n".join(parts)


def cvs_stamp(master: Path) -> tuple[tuple[str, int], ...]:
    """Changes whenever any CV is added, removed or edited."""
    folder = cv_dir(master)
    return tuple(sorted((p.name, p.stat().st_mtime_ns) for p in folder.glob("*.md")))


def cv_path(master: Path, name: str) -> Path:
    """The file for a CV name, refusing anything that isn't a plain stored name."""
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", name):
        raise CvError(f"Invalid CV name: {name!r}")
    return cv_dir(master) / f"{name}.md"


def free_name(master: Path, wanted: str) -> str:
    """A new, unused CV name based on `wanted` (e.g. an upload's file name)."""
    base = slugify(Path(wanted).stem) or "cv"
    if base == master.stem:
        base = f"{base}-copy"
    name, n = base, 2
    while cv_path(master, name).exists():
        name, n = f"{base}-{n}", n + 1
    return name


def save_cv(master: Path, name: str, text: str) -> Path:
    path = cv_path(master, name)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_normalise(text))
    return path


def delete_cv(master: Path, name: str) -> None:
    path = cv_path(master, name)
    if path.resolve() == master.resolve():
        raise CvError("The master CV can't be deleted; make another CV the master first.")
    path.unlink(missing_ok=True)


def make_master(master: Path, name: str, backup_dir: Path) -> Path | None:
    """Copy a reference CV's text into the master file. Returns the backup of the old
    master, if there was one. Changing the master re-ranks every job (cache key)."""
    source = cv_path(master, name)
    if not source.is_file():
        raise CvError(f"No CV named {name!r}")
    if source.resolve() == master.resolve():
        return None
    backup = None
    if master.is_file():
        backup_dir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        backup = backup_dir / f"{master.stem}-{stamp}.md"
        shutil.copy2(master, backup)
    master.parent.mkdir(parents=True, exist_ok=True)
    master.write_text(source.read_text())
    return backup


# --- conversion ---------------------------------------------------------------


def convert_upload(filename: str, data: bytes) -> str:
    """Markdown text from an uploaded CV (PDF, DOCX, Markdown or plain text)."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise CvError(f"File is larger than {MAX_UPLOAD_BYTES // 1024 // 1024} MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = pdf_to_markdown(data)
    elif suffix == ".docx":
        text = docx_to_markdown(data)
    elif suffix in (".md", ".markdown", ".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
    else:
        raise CvError(f"Unsupported file type {suffix or '(none)'}: use {', '.join(UPLOAD_TYPES)}")
    text = _normalise(text)
    if not text.strip():
        raise CvError("No text found in the file (a scanned PDF needs OCR first)")
    return text


def pdf_to_markdown(data: bytes) -> str:
    """Plain text from a PDF. PDFs carry no headings, so review the result."""
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        raise CvError(f"Couldn't read the PDF: {exc}") from exc
    lines = []
    for line in "\n".join(pages).splitlines():
        stripped = line.strip()
        # PDF bullets come through as symbols; make them Markdown list items.
        stripped = re.sub(r"^[●•▪◦‣∙·]\s*", "- ", stripped)
        lines.append(stripped)
    return "\n".join(lines)


def docx_to_markdown(data: bytes) -> str:
    """Markdown from a Word document: headings, list items and table rows."""
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise CvError(f"Couldn't read the Word document: {exc}") from exc
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        style = (paragraph.style.name if paragraph.style is not None else "").lower()
        if not text:
            lines.append("")
        elif style == "title":
            lines.append(f"# {text}")
        elif style.startswith("heading"):
            level = int(re.sub(r"\D", "", style) or 1)
            lines.append(f"{'#' * min(level + 1, 6)} {text}")
        elif "list" in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(c for c in cells if c))
    return "\n".join(lines)


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"
