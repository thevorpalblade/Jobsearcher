"""Markdown to Word (python-docx) and PDF (headless LibreOffice, from the same .docx,
so the two always match). Without LibreOffice only Word and Markdown are produced."""

from __future__ import annotations

import logging
import re
import shutil
import subprocess
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

log = logging.getLogger(__name__)

FONT = "Arial"  # LibreOffice maps it to Liberation Sans; both have åäö
SIGNATURE_WIDTH = Cm(4.5)
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")
_SIGN_HERE = "\x00signature"  # where the letter's signature image goes


def _runs(
    paragraph, text: str, size: float | None = None, bold: bool = False, font: str = FONT
) -> None:
    """Add `text` to a paragraph, turning **bold**, *italic* and `code` into runs."""
    for part in _INLINE.split(text):
        if not part:
            continue
        run_bold, italic = bold, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, run_bold = part[2:-2], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        elif part.startswith("`") and part.endswith("`"):
            part = part[1:-1]
        run = paragraph.add_run(part)
        run.bold, run.italic = run_bold, italic
        run.font.name = font
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        if size:
            run.font.size = Pt(size)


def _rule_under(paragraph) -> None:
    """A thin line under a section heading."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "808080")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _setup(document, size: float, font: str = FONT) -> None:
    for section in document.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.paragraph_format.space_after = Pt(4)


def markdown_to_docx(
    markdown: str, kind: str, path: Path, font: str = FONT, signature: Path | None = None
) -> None:
    """Write `markdown` to `path` as a Word file. `kind` is "cv" or "letter". A
    `signature` image goes above the letter's last line (the signed name)."""
    document = docx.Document()
    cv = kind == "cv"
    _setup(document, 10 if cv else 11, font)
    lines_in = markdown.replace("\r\n", "\n").split("\n")
    sign_at = None
    if signature is not None and not cv:
        sign_at = max((i for i, line in enumerate(lines_in) if line.strip()), default=None)
    after_title = False
    paragraph_lines: list[str] = []

    def flush() -> None:
        nonlocal after_title
        if not paragraph_lines:
            return
        lines = [line.strip() for line in paragraph_lines]
        paragraph_lines.clear()
        p = document.add_paragraph()
        if after_title:  # the contact line under the name
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            after_title, size = False, 9.5
        else:
            if not cv:
                p.paragraph_format.space_after = Pt(9)
            size = None
        for i, line in enumerate(lines):  # a single newline stays a line break (sign-offs)
            if i:
                p.add_run().add_break()
            if line == _SIGN_HERE:
                p.add_run().add_picture(str(signature), width=SIGNATURE_WIDTH)
                continue
            _runs(p, line, size=size, font=font)

    for n, raw in enumerate(lines_in):
        if n == sign_at:
            paragraph_lines.append(_SIGN_HERE)
        line = raw.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        bullet = _BULLET.match(line)
        if not line.strip():
            flush()
        elif heading:
            flush()
            level, text = len(heading.group(1)), heading.group(2).strip()
            p = document.add_paragraph()
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _runs(p, text, size=17 if cv else 14, bold=True, font=font)
                p.paragraph_format.space_after = Pt(2)
                after_title = True
            elif level == 2:
                p.paragraph_format.space_before = Pt(9)
                p.paragraph_format.space_after = Pt(3)
                _runs(p, text.upper(), size=10.5, bold=True, font=font)
                _rule_under(p)
            else:
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(1)
                _runs(p, text, size=10.5, bold=True, font=font)
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        elif bullet:
            flush()
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            _runs(p, bullet.group(1).strip(), font=font)
        elif re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            flush()
        else:
            paragraph_lines.append(line)
    flush()
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def docx_to_pdf(docx_path: Path, timeout_s: float = 120) -> Path | None:
    """The PDF next to `docx_path`, or None when LibreOffice isn't available or fails."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        return None
    try:
        subprocess.run(
            [
                soffice,
                "--headless",
                # A private profile: a running LibreOffice would otherwise swallow the call.
                "-env:UserInstallation=file:///tmp/jobsearcher-libreoffice",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx_path.parent),
                str(docx_path),
            ],
            capture_output=True,
            timeout=timeout_s,
            check=True,
        )
    except (subprocess.SubprocessError, OSError) as exc:
        log.warning("PDF conversion of %s failed: %s", docx_path.name, exc)
        return None
    pdf = docx_path.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def render_files(
    folder: Path,
    cv_markdown: str,
    letter_markdown: str,
    font: str = FONT,
    signature: Path | None = None,
) -> list[str]:
    """Write cv/letter as .md, .docx and (when possible) .pdf; return the file names.
    `font` and `signature` are the letter's; the CV keeps the default look."""
    folder.mkdir(parents=True, exist_ok=True)
    names: list[str] = []
    for stem, kind, text in (("cv", "cv", cv_markdown), ("letter", "letter", letter_markdown)):
        (folder / f"{stem}.md").write_text(text.strip() + "\n")
        if kind == "letter":
            markdown_to_docx(text, kind, folder / f"{stem}.docx", font, signature)
        else:
            markdown_to_docx(text, kind, folder / f"{stem}.docx")
        names += [f"{stem}.md", f"{stem}.docx"]
        if docx_to_pdf(folder / f"{stem}.docx"):
            names.append(f"{stem}.pdf")
    return names
