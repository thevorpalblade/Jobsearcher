import shutil
import threading
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from conftest import make_assessment, make_job

from jobsearcher.companies.config import Company
from jobsearcher.config import Config
from jobsearcher.drafting.core import (
    Claim,
    DraftContent,
    GroundingCheck,
    Request,
    draft_hash,
    generate_draft,
    unsupported_numbers,
)
from jobsearcher.drafting.manager import DraftManager
from jobsearcher.drafting.render import markdown_to_docx, render_files
from jobsearcher.drafting.service import (
    DraftError,
    Llms,
    choose_contact,
    draft_company,
    draft_job,
    job_request,
    load_cvs,
)
from jobsearcher.llm import LLMError, LLMResult, LLMUsage
from jobsearcher.models import Contact
from jobsearcher.ranking.ranker import Ranking
from jobsearcher.store import Store

CV = "# Alex Example\n\n## Experience\n\n- Led the rollout of 3 new warehouses in 2023, growing from 5 to 8 locations.\n"
OTHER = "# Older CV\n\n- Directed HR and IT for a multi-site organisation.\n"


def content(letter="Dear Hiring Manager,\n\nI led the rollout of 3 new warehouses.\n\nAlex", cv=CV):
    return DraftContent(cover_letter=letter, cv=cv, notes=["Emphasised the integration work."])


def verdict(*claims):
    return GroundingCheck(
        claims=[Claim(claim=c, supported=ok, evidence="x" if ok else None) for c, ok in claims]
    )


class FakeLLM:
    """Answers each call with the next prepared output; records every prompt."""

    def __init__(self, model, *outputs):
        self.model, self.outputs, self.calls = model, list(outputs), []

    def complete(self, *, system, prompt, context="", schema=None):
        self.calls.append({"system": system, "context": context, "prompt": prompt})
        out = self.outputs.pop(0)
        if isinstance(out, Exception):
            raise out
        usage = LLMUsage(model=self.model, input_tokens=1, output_tokens=1, billed=False)
        return LLMResult(text="", usage=usage, parsed=out)


def request(**kwargs):
    base = {"key": "job1", "prompt": "Write.", "identity": ["ad v1"], "sources": ["Stockholm 2027"]}
    return Request(**(base | kwargs))


def fake_render(folder, cv, letter):
    return ["cv.md", "letter.md"]


def run(store, drafter, checker, req=None, cvs=None, **kwargs):
    return generate_draft(
        store,
        req or request(),
        cvs or [("master", CV)],
        drafter,
        checker,
        Path("drafts"),
        fake_render,
        **kwargs,
    )


# --- number guard -------------------------------------------------------------------


def test_unsupported_numbers():
    sources = ["Led 3 new warehouses in 2023, 5 to 8 locations; 10+ years; 25% savings"]
    assert unsupported_numbers("In 2023 I added 3 warehouses; 10+ years; 25% saved", sources) == []
    flagged = unsupported_numbers("I cut costs by 30% in 2019 over 12 years", sources)
    assert [c.claim for c in flagged] == [
        "The figure “12” appears in the draft",
        "The figure “2019” appears in the draft",
        "The figure “30%” appears in the draft",
    ]
    assert all(not c.supported for c in flagged)


# --- generation, cache, repair ---------------------------------------------------------


def test_draft_is_checked_stored_and_cached():
    store = Store(":memory:")
    drafter = FakeLLM("opus", content())
    checker = FakeLLM("glm", verdict(("Led the rollout of 3 new warehouses", True)))
    draft = run(store, drafter, checker)
    assert not draft.needs_review and (draft.claims_checked, draft.repaired) == (1, False)
    assert draft.files == ["cv.md", "letter.md"] and draft.base_cv == "master"
    assert (draft.model, draft.check_model) == ("opus", "glm")
    assert "Base CV (master)" in drafter.calls[0]["context"]
    assert "never invent" in drafter.calls[0]["system"].lower()
    assert "verify" in checker.calls[0]["system"].lower()

    again = run(store, drafter, checker)  # nothing changed: no model calls at all
    assert again.input_hash == draft.input_hash and len(drafter.calls) == 1
    assert len(store.list_drafts("job1")) == 1

    drafter.outputs, checker.outputs = [content()], [verdict()]
    run(store, drafter, checker, request(instructions="Lead with M&A"))  # new instructions
    assert len(drafter.calls) == 2 and len(store.list_drafts("job1")) == 2
    assert draft_hash(request(), [("master", CV)], "opus") != draft_hash(
        request(), [("master", CV + "x")], "opus"
    )
    assert draft_hash(request(), [("master", CV)], "opus") != draft_hash(
        request(identity=["ad v2"]), [("master", CV)], "opus"
    )


def test_unsupported_claims_get_one_repair_round():
    store = Store(":memory:")
    first = content(letter="Dear Hiring Manager,\n\nI ran a team of 40 engineers.\n\nAlex")
    fixed = content(letter="Dear Hiring Manager,\n\nI led the rollout of 3 new warehouses.\n\nAlex")
    drafter = FakeLLM("opus", first, fixed)
    checker = FakeLLM(
        "glm",
        verdict(("Ran a team of 40 engineers", False)),
        verdict(("Led the rollout of 3 new warehouses", True)),
    )
    draft = run(store, drafter, checker)
    assert draft.repaired and not draft.needs_review and "40 engineers" not in draft.cover_letter
    repair = drafter.calls[1]["prompt"]
    assert "NOT supported" in repair and "Ran a team of 40 engineers" in repair
    assert "I ran a team of 40 engineers" in repair  # the previous draft is shown to fix


def test_still_unsupported_after_repair_needs_review():
    store = Store(":memory:")
    bad = content(letter="Dear Hiring Manager,\n\nI hold a PMP certification.\n\nAlex")
    drafter, checker = (
        FakeLLM("opus", bad, bad),
        FakeLLM("glm", verdict(("Holds a PMP", False)), verdict(("Holds a PMP", False))),
    )
    draft = run(store, drafter, checker)
    assert draft.needs_review and [c.claim for c in draft.flagged] == ["Holds a PMP"]
    assert draft.repaired and len(drafter.calls) == 2  # one repair, no more
    assert store.get_draft("job1", draft.input_hash) is not None  # saved, marked, not hidden


def test_invented_figures_are_caught_even_when_the_model_check_passes():
    store = Store(":memory:")
    letter = "Dear Hiring Manager,\n\nI grew revenue by 45% over 12 years.\n\nAlex"
    drafter = FakeLLM("opus", content(letter=letter), content(letter=letter))
    checker = FakeLLM("glm", verdict(("Grew revenue", True)), verdict(("Grew revenue", True)))
    draft = run(store, drafter, checker)
    assert draft.needs_review
    assert {c.claim for c in draft.flagged} == {
        "The figure “12” appears in the draft",
        "The figure “45%” appears in the draft",
    }
    # Figures from the ad (a source) are fine in the letter, not in the CV.
    ok = content(letter="Dear Hiring Manager,\n\nI'm excited about 2027.\n\nAlex")
    drafter2, checker2 = FakeLLM("opus", ok), FakeLLM("glm", verdict())
    assert not run(store, drafter2, checker2, request(key="job2")).needs_review


def test_a_failed_check_marks_the_draft_without_repairing():
    store = Store(":memory:")
    drafter = FakeLLM("opus", content())
    checker = FakeLLM("glm", LLMError("504"), LLMError("504"))
    draft = run(store, drafter, checker)
    assert draft.needs_review and "grounding check couldn't run" in draft.check_error
    assert not draft.repaired and len(drafter.calls) == 1 and draft.flagged == []


# --- contacts and requests ----------------------------------------------------------------


def test_contact_choice_and_requests():
    named = Contact(name="Anna Svensson", role="HR-chef", provenance="llm:ad_text")
    official = Contact(name="Per Persson", provenance="platsbanken:application_contacts")
    generic = Contact(email="jobb@acme.se", role="generic mailbox", provenance="x")
    assert choose_contact(make_job(1, "J", contacts=[generic])) is None
    assert choose_contact(make_job(1, "J", contacts=[named, official])).name == "Per Persson"
    assert choose_contact(make_job(1, "J", contacts=[generic, named])).name == "Anna Svensson"

    job = make_job(1, "HR Business Partner", contacts=[named])
    ranking = Ranking(
        job_id=job.id, input_hash="h", model="m", assessment=make_assessment(80, 70)
    ).model_dump_json()
    req = job_request(job, ranking, "Lead with M&A")
    assert req.addressed_to == "Anna Svensson" and req.key == job.id
    assert "Dear Anna Svensson," in req.prompt and "HR-chef" in req.prompt
    assert "Her instructions: Lead with M&A" in req.prompt and "HR partnering" in req.prompt
    anon = job_request(make_job(2, "J"), None)
    assert "Dear Hiring Manager," in anon.prompt and "don't invent one" in anon.prompt
    assert job_request(job, "not json", "").key == job.id  # an old ranking is just skipped


# --- the service with a real store and files --------------------------------------------------


@pytest.fixture
def env(tmp_path):
    (tmp_path / "cvs").mkdir()
    (tmp_path / "cvs" / "master.md").write_text(CV)
    (tmp_path / "cvs" / "older.md").write_text(OTHER)
    (tmp_path / "ranking.yaml").write_text("target_roles:\n  - name: HR Business Partner\n")
    config = Config(
        data_dir=tmp_path / "data", cv_path=tmp_path / "cvs" / "master.md",
        ranking_config=tmp_path / "ranking.yaml",
    )  # fmt: skip
    store = Store(config.db_path)
    return config, store


def test_load_cvs_orders_the_base_first(env):
    config, _ = env
    assert [n for n, _ in load_cvs(config)] == ["master", "older"]
    assert [n for n, _ in load_cvs(config, "older")] == ["older", "master"]
    assert [n for n, _ in load_cvs(config, "nope")] == ["master", "older"]
    config.cv_path.unlink()
    (config.cv_path.parent / "older.md").unlink()
    with pytest.raises(DraftError, match="No CV"):
        load_cvs(config)


def test_draft_job_end_to_end(env):
    config, store = env
    job = make_job(1, "HR Business Partner", company="Acme")
    store.upsert_job(job)
    llms = Llms(FakeLLM("opus", content()), FakeLLM("glm", verdict(("Led 3 warehouses", True))))
    draft = draft_job(
        config, store, job.id, "Be brief", base_cv="older", llms=llms, render=fake_render
    )
    assert draft.base_cv == "older" and draft.instructions == "Be brief"
    assert "Other CV: master" in llms.draft.calls[0]["context"]
    with pytest.raises(DraftError, match="No job"):
        draft_job(config, store, "nope", llms=llms)


def test_spontaneous_drafts_need_a_news_reason(env):
    config, store = env
    acme = Company(name="Acme AB")
    with pytest.raises(DraftError, match="No recent news"):
        draft_company(
            config, store, acme, llms=Llms(FakeLLM("o"), FakeLLM("g")), render=fake_render
        )

    now = datetime.now(UTC)
    store.save_news_items(
        [{"id": "n1", "company": acme.slug, "title": "Acme buys Beta", "url": "https://n/1",
          "domain": "di.se", "published_at": (now - timedelta(days=3)).isoformat(),
          "fetched_at": now.isoformat()}]
    )  # fmt: skip
    store.save_signal("n1", "merger_acquisition", 85, "Acme is acquiring Beta.", "glm", "1")
    llms = Llms(FakeLLM("opus", content()), FakeLLM("glm", verdict()))
    draft = draft_company(config, store, acme, llms=llms, render=fake_render)
    assert draft.key == "company:acme-ab"
    prompt = llms.draft.calls[0]["prompt"]
    assert "SPONTANEOUS" in prompt and "Acme is acquiring Beta." in prompt
    assert "HR Business Partner" in prompt  # her target roles
    assert "Dear Hiring Manager," in prompt


# --- rendering -------------------------------------------------------------------------------


def test_markdown_to_docx_structure(tmp_path):
    import docx

    letter = "Dear Anna,\n\nI led **3 warehouses** and *grew* them.\n\n- first point\n- second point\n\nBest regards,\nAlex"
    markdown_to_docx(letter, "letter", tmp_path / "l.docx")
    document = docx.Document(str(tmp_path / "l.docx"))
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs]
    assert paragraphs[0] == ("Normal", "Dear Anna,")
    assert ("List Bullet", "first point") in paragraphs
    assert paragraphs[-1][1] == "Best regards,\nAlex"  # the sign-off keeps its line break
    bold = [r.text for r in document.paragraphs[1].runs if r.bold]
    assert bold == ["3 warehouses"]

    markdown_to_docx(CV, "cv", tmp_path / "c.docx")
    texts = [p.text for p in docx.Document(str(tmp_path / "c.docx")).paragraphs]
    assert texts[:2] == ["Alex Example", "EXPERIENCE"]  # name; sections in capitals


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_render_files_makes_word_and_pdf(tmp_path):
    from pypdf import PdfReader

    names = render_files(tmp_path / "v1", CV, "Hej — åäö Södertälje.\n\nAlex")
    assert names == ["cv.md", "cv.docx", "cv.pdf", "letter.md", "letter.docx", "letter.pdf"]
    text = PdfReader(str(tmp_path / "v1" / "letter.pdf")).pages[0].extract_text()
    assert "åäö Södertälje" in text


def test_render_without_libreoffice_gives_word_and_markdown(tmp_path, monkeypatch):
    monkeypatch.setattr("jobsearcher.drafting.render.shutil.which", lambda name: None)
    assert render_files(tmp_path, CV, "Hi") == ["cv.md", "cv.docx", "letter.md", "letter.docx"]


# --- the background runner ------------------------------------------------------------------


def test_manager_runs_one_draft_at_a_time_and_reports_failures(tmp_path):
    Store(tmp_path / "m.db").close()
    manager = DraftManager(lambda: Config(), tmp_path / "m.db")
    gate, order = threading.Event(), []

    def slow(config, store):
        order.append("slow")
        gate.wait(5)

    def fails(config, store):
        raise LLMError("usage limit reached")

    first = manager.submit("a", slow)
    assert manager.submit("a", slow) is first  # already queued/running: not queued twice
    second = manager.submit("b", fails)
    time.sleep(0.2)
    assert (first.state, second.state, manager.pending()) == ("running", "queued", 2)
    gate.set()
    deadline = time.time() + 5
    while manager.pending() and time.time() < deadline:
        time.sleep(0.02)
    assert first.state == "done" and second.state == "failed"
    assert second.error == "usage limit reached" and order == ["slow"]
    assert manager.submit("a", lambda c, s: None) is not first  # a finished key can run again


def test_a_slow_primary_check_falls_back_to_the_next_model():
    store = Store(":memory:")
    bad = content(letter="Dear Hiring Manager,\n\nI hold a PMP.\n\nAlex")
    good = content(letter="Dear Hiring Manager,\n\nI led the rollout of 3 new warehouses.\n\nAlex")
    drafter = FakeLLM("opus", bad, good)
    primary = FakeLLM("glm", LLMError("504 gateway timeout"))  # GLM is queued behind other work
    fallback = FakeLLM(
        "haiku", verdict(("Holds a PMP", False)), verdict(("Led 3 warehouses", True))
    )
    draft = run(store, drafter, primary, fallback_llm=fallback)
    assert draft.check_model == "haiku" and draft.check_error is None
    assert draft.repaired and not draft.needs_review  # the fallback's verdict drove the repair
    assert len(primary.calls) == 1  # asked first, never again once it failed


def test_both_checkers_down_marks_the_draft_for_review():
    store = Store(":memory:")
    draft = run(
        store,
        FakeLLM("opus", content()),
        FakeLLM("glm", LLMError("504")),
        fallback_llm=FakeLLM("haiku", LLMError("usage limit")),
    )
    assert draft.needs_review and draft.check_model is None
    assert "usage limit" in draft.check_error  # the last failure is the one shown
